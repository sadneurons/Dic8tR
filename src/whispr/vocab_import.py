"""Vocabulary import dialog for Whispr.

Allows importing corrections and expansions from Excel (.xlsx),
CSV, or JSON files into the user's vocabulary.json.
"""

import csv
import json
import logging
from pathlib import Path

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from whispr.config import DEFAULT_VOCAB_PATH, load_vocabulary

logger = logging.getLogger(__name__)

SUPPORTED_FILTERS = "All supported (*.xlsx *.csv *.json);;Excel (*.xlsx);;CSV (*.csv);;JSON (*.json)"


class VocabImportDialog(QDialog):
    """Dialog for importing vocabulary from external files.

    Supports:
    - Excel (.xlsx): sheet "corrections" (col A=spoken, B=replacement)
                     and/or sheet "expansions" (col A=spoken, B=replacement)
                     OR single sheet with columns: type, spoken, replacement
    - CSV: columns type,spoken,replacement
    - JSON: same format as vocabulary.json
    """

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("Import Vocabulary")
        self.setMinimumSize(580, 450)

        self._corrections: dict[str, str] = {}
        self._expansions: dict[str, str] = {}

        layout = QVBoxLayout(self)

        # File picker
        file_row = QHBoxLayout()
        self._file_label = QLabel("No file selected")
        self._file_label.setStyleSheet("color: #777;")
        file_row.addWidget(self._file_label, 1)
        browse_btn = QPushButton("Browse...")
        browse_btn.clicked.connect(self._browse_file)
        file_row.addWidget(browse_btn)
        layout.addLayout(file_row)

        # Import mode
        mode_row = QHBoxLayout()
        mode_row.addWidget(QLabel("Mode:"))
        self._mode_combo = QComboBox()
        self._mode_combo.addItems(["Merge (add new, keep existing)", "Replace all"])
        mode_row.addWidget(self._mode_combo)
        mode_row.addStretch()
        layout.addLayout(mode_row)

        # Preview tables
        self._corrections_group = QGroupBox("Corrections (0)")
        corrections_layout = QVBoxLayout(self._corrections_group)
        self._corrections_table = QTableWidget(0, 2)
        self._corrections_table.setHorizontalHeaderLabels(["Spoken Form", "Replacement"])
        self._corrections_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        corrections_layout.addWidget(self._corrections_table)
        layout.addWidget(self._corrections_group)

        self._expansions_group = QGroupBox("Expansions (0)")
        expansions_layout = QVBoxLayout(self._expansions_group)
        self._expansions_table = QTableWidget(0, 2)
        self._expansions_table.setHorizontalHeaderLabels(["Spoken Form", "Expansion Text"])
        self._expansions_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        expansions_layout.addWidget(self._expansions_table)
        layout.addWidget(self._expansions_group)

        # Status
        self._status = QLabel("")
        self._status.setStyleSheet("font-size: 11px;")
        layout.addWidget(self._status)

        # Buttons
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        self._ok_btn = buttons.button(QDialogButtonBox.StandardButton.Ok)
        self._ok_btn.setText("Import")
        self._ok_btn.setEnabled(False)
        buttons.accepted.connect(self._on_import)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _browse_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "Select Vocabulary File", "", SUPPORTED_FILTERS
        )
        if not path:
            return

        self._file_label.setText(path)
        self._file_label.setStyleSheet("")
        self._parse_file(Path(path))

    def _parse_file(self, path: Path) -> None:
        self._corrections = {}
        self._expansions = {}

        try:
            suffix = path.suffix.lower()
            if suffix == ".json":
                self._load_json(path)
            elif suffix == ".csv":
                self._load_csv(path)
            elif suffix == ".xlsx":
                self._load_excel(path)
            else:
                self._set_error(f"Unsupported format: {suffix}")
                return
        except Exception as e:
            self._set_error(f"Error reading file: {e}")
            logger.error("Vocab import parse error: %s", e)
            return

        self._update_preview()
        total = len(self._corrections) + len(self._expansions)
        self._ok_btn.setEnabled(total > 0)
        if total == 0:
            self._set_status("No entries found in file.", error=True)
        else:
            self._set_status(f"Found {len(self._corrections)} corrections, {len(self._expansions)} expansions.")

    def _load_json(self, path: Path) -> None:
        with open(path) as f:
            data = json.load(f)
        self._corrections = data.get("corrections", {})
        self._expansions = data.get("expansions", {})

    def _load_csv(self, path: Path) -> None:
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                entry_type = row.get("type", "").strip().lower()
                spoken = row.get("spoken", "").strip()
                replacement = row.get("replacement", "").strip()
                if not spoken or not replacement:
                    continue
                if entry_type == "correction":
                    self._corrections[spoken] = replacement
                elif entry_type == "expansion":
                    self._expansions[spoken] = replacement

    def _load_excel(self, path: Path) -> None:
        try:
            import openpyxl
        except ImportError:
            self._set_error(
                "openpyxl is required for Excel import.\n"
                "Install with: pip install openpyxl"
            )
            return

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet_names = [s.lower() for s in wb.sheetnames]

        # Try named sheets first
        if "corrections" in sheet_names:
            idx = sheet_names.index("corrections")
            ws = wb[wb.sheetnames[idx]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 2 and row[0] and row[1]:
                    self._corrections[str(row[0]).strip()] = str(row[1]).strip()

        if "expansions" in sheet_names:
            idx = sheet_names.index("expansions")
            ws = wb[wb.sheetnames[idx]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 2 and row[0] and row[1]:
                    self._expansions[str(row[0]).strip()] = str(row[1]).strip()

        # If no named sheets, try first sheet with type,spoken,replacement columns
        if not self._corrections and not self._expansions:
            ws = wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 3 and row[0] and row[1] and row[2]:
                    entry_type = str(row[0]).strip().lower()
                    spoken = str(row[1]).strip()
                    replacement = str(row[2]).strip()
                    if entry_type == "correction":
                        self._corrections[spoken] = replacement
                    elif entry_type == "expansion":
                        self._expansions[spoken] = replacement

        wb.close()

    def _update_preview(self) -> None:
        self._corrections_group.setTitle(f"Corrections ({len(self._corrections)})")
        self._corrections_table.setRowCount(len(self._corrections))
        for i, (spoken, replacement) in enumerate(self._corrections.items()):
            self._corrections_table.setItem(i, 0, QTableWidgetItem(spoken))
            self._corrections_table.setItem(i, 1, QTableWidgetItem(replacement))

        self._expansions_group.setTitle(f"Expansions ({len(self._expansions)})")
        self._expansions_table.setRowCount(len(self._expansions))
        for i, (spoken, expansion) in enumerate(self._expansions.items()):
            self._expansions_table.setItem(i, 0, QTableWidgetItem(spoken))
            self._expansions_table.setItem(i, 1, QTableWidgetItem(expansion))

    def _on_import(self) -> None:
        """Merge or replace the current vocabulary with imported entries."""
        replace_mode = self._mode_combo.currentIndex() == 1

        vocab = load_vocabulary()

        if replace_mode:
            vocab["corrections"] = self._corrections
            vocab["expansions"] = self._expansions
        else:
            # Merge: new entries added, existing preserved
            for k, v in self._corrections.items():
                if k not in vocab.get("corrections", {}):
                    vocab.setdefault("corrections", {})[k] = v
            for k, v in self._expansions.items():
                if k not in vocab.get("expansions", {}):
                    vocab.setdefault("expansions", {})[k] = v

        try:
            with open(DEFAULT_VOCAB_PATH, "w") as f:
                json.dump(vocab, f, indent=2, ensure_ascii=False)
            logger.info(
                "Vocabulary imported: %d corrections, %d expansions (%s mode)",
                len(self._corrections), len(self._expansions),
                "replace" if replace_mode else "merge",
            )
            self.accept()
        except OSError as e:
            self._set_error(f"Failed to save vocabulary: {e}")

    def _set_status(self, msg: str, error: bool = False) -> None:
        self._status.setText(msg)
        self._status.setStyleSheet(
            f"font-size: 11px; color: {'#C62828' if error else '#555'};"
        )

    def _set_error(self, msg: str) -> None:
        self._set_status(msg, error=True)
        self._ok_btn.setEnabled(False)
